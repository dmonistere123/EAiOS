#!/usr/bin/env python3
"""EAiOS Knowledge sidecar — governed RAG source indexing on loopback.

Phase 5.2, path A: upload/URL -> extract -> chunk -> SQLite FTS5.
Vectors later, behind the same HTTP surface. Loopback-only bind; the app
reaches it through the vite /knowledge-api proxy, so the browser never
talks to it directly (same trust pattern as the Composio proxy).

Endpoints (JSON in/out unless noted):
  GET    /health                     -> {"ok": true, "sources": N, "chunks": M}
  GET    /sources                    -> {"sources": [...]}
  POST   /sources/upload             multipart: file + name/scope/citationEnabled/allowedAgentIds
  POST   /sources/url                {"url", "name"?, "scope", "citationEnabled", "allowedAgentIds"?}
  POST   /sources/<id>/reindex
  DELETE /sources/<id>
  GET    /search?q=...&limit=8       -> {"results": [{sourceId, sourceName, chunkIndex, snippet, score}]}
  GET    /chunks/<chunk_id>          -> full chunk text + source metadata (drill-down)

Scope enforcement (spec: "scope is enforced by the runtime"): pass agent_id
on /search and private sources are withheld; agent-scoped sources only
surface to agents named in allowedAgentIds. No agent_id = executive context
(sees everything). citationEnabled flags whether answers may QUOTE a chunk.
"""
import json
import os
import re
import sqlite3
import sys
import time
import traceback
import urllib.request
import uuid
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from html.parser import HTMLParser
from pathlib import Path

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
FILES = DATA / "files"
DB = DATA / "knowledge.db"
HOST, PORT = "127.0.0.1", 9121
CHUNK_SIZE = 1000
CHUNK_OVERLAP = 150
MAX_UPLOAD = 50 * 1024 * 1024

SCHEMA = """
CREATE TABLE IF NOT EXISTS sources (
  id TEXT PRIMARY KEY,
  type TEXT NOT NULL,             -- file | url | text
  name TEXT NOT NULL,
  uri TEXT,
  stored_path TEXT,
  scope TEXT NOT NULL,            -- private | workspace | agent
  allowed_agent_ids TEXT,         -- JSON array or NULL
  citation_enabled INTEGER NOT NULL,
  indexing_status TEXT NOT NULL,  -- pending | processing | ready | failed | stale
  error TEXT,
  freshness_at TEXT,
  created_at TEXT NOT NULL
);
CREATE VIRTUAL TABLE IF NOT EXISTS chunks USING fts5(
  chunk_id UNINDEXED, source_id UNINDEXED, chunk_index UNINDEXED, text
);
"""


def now_iso() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


def db() -> sqlite3.Connection:
    conn = sqlite3.connect(DB)
    conn.row_factory = sqlite3.Row
    conn.executescript(SCHEMA)
    return conn


# ---------- extraction ----------

def extract_pdf(path: Path) -> str:
    import fitz  # pymupdf

    with fitz.open(path) as doc:
        return "\n\n".join(page.get_text() for page in doc)


def extract_docx(path: Path) -> str:
    import docx

    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs if p.text.strip())


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    prs = Presentation(path)
    out = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                out.append(shape.text_frame.text)
    return "\n\n".join(out)


class _HtmlText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts: list[str] = []
        self._skip = 0

    def handle_starttag(self, tag, attrs):
        if tag in ("script", "style", "noscript"):
            self._skip += 1
        if tag in ("p", "br", "div", "li", "h1", "h2", "h3", "tr"):
            self.parts.append("\n")

    def handle_endtag(self, tag):
        if tag in ("script", "style", "noscript") and self._skip:
            self._skip -= 1

    def handle_data(self, data):
        if not self._skip:
            self.parts.append(data)


def extract_html(text: str) -> str:
    p = _HtmlText()
    p.feed(text)
    out = re.sub(r"\n\s*\n+", "\n\n", "".join(p.parts))
    return re.sub(r"[ \t]+", " ", out).strip()


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".docx":
        return extract_docx(path)
    if ext == ".pptx":
        return extract_pptx(path)
    raw = path.read_bytes()
    text = raw.decode("utf-8", errors="replace")
    if ext in (".html", ".htm"):
        return extract_html(text)
    return text  # .txt .md .csv .json and friends index fine as-is


# ---------- chunking + indexing ----------

def chunk_text(text: str) -> list[str]:
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        return []
    chunks = []
    start = 0
    while start < len(text):
        end = min(start + CHUNK_SIZE, len(text))
        # prefer breaking at a paragraph/sentence boundary within the window
        if end < len(text):
            cut = max(text.rfind("\n\n", start, end), text.rfind(". ", start, end))
            if cut > start + CHUNK_SIZE // 2:
                end = cut + 1
        chunks.append(text[start:end].strip())
        start = end - CHUNK_OVERLAP if end < len(text) else end
    return [c for c in chunks if c]


def index_source(conn: sqlite3.Connection, source_id: str) -> tuple[int, str | None]:
    """(Re)build chunks for a source. Returns (chunk_count, error)."""
    row = conn.execute("SELECT * FROM sources WHERE id = ?", (source_id,)).fetchone()
    if not row:
        return 0, "source not found"
    conn.execute("UPDATE sources SET indexing_status = 'processing', error = NULL WHERE id = ?", (source_id,))
    conn.commit()
    try:
        text = extract_text(Path(row["stored_path"]))
        chunks = chunk_text(text)
        if not chunks:
            raise ValueError("no extractable text")
        conn.execute("DELETE FROM chunks WHERE source_id = ?", (source_id,))
        conn.executemany(
            "INSERT INTO chunks (chunk_id, source_id, chunk_index, text) VALUES (?,?,?,?)",
            [(f"{source_id}:{i}", source_id, i, c) for i, c in enumerate(chunks)],
        )
        conn.execute(
            "UPDATE sources SET indexing_status = 'ready', freshness_at = ?, error = NULL WHERE id = ?",
            (now_iso(), source_id),
        )
        conn.commit()
        return len(chunks), None
    except Exception as e:  # extraction failure is data, not a crash
        conn.execute(
            "UPDATE sources SET indexing_status = 'failed', error = ? WHERE id = ?",
            (f"{type(e).__name__}: {e}"[:300], source_id),
        )
        conn.commit()
        return 0, str(e)


def fetch_url(url: str) -> tuple[bytes, str]:
    req = urllib.request.Request(url, headers={"User-Agent": "EAiOS-Knowledge/0.1"})
    with urllib.request.urlopen(req, timeout=30) as r:
        body = r.read(MAX_UPLOAD)
        ctype = r.headers.get_content_type()
    return body, ctype


def source_row(r: sqlite3.Row) -> dict:
    return {
        "id": r["id"],
        "type": r["type"],
        "name": r["name"],
        "uri": r["uri"],
        "scope": r["scope"],
        "allowedAgentIds": json.loads(r["allowed_agent_ids"]) if r["allowed_agent_ids"] else None,
        "indexingStatus": r["indexing_status"],
        "error": r["error"],
        "freshnessAt": r["freshness_at"],
        "citationEnabled": bool(r["citation_enabled"]),
    }


# ---------- minimal multipart parsing (stdlib) ----------

def parse_multipart(body: bytes, boundary: str) -> dict[str, tuple[str | None, bytes]]:
    """Returns {field: (filename_or_None, content)}."""
    out: dict[str, tuple[str | None, bytes]] = {}
    delim = b"--" + boundary.encode()
    for part in body.split(delim):
        part = part.strip(b"\r\n")
        if not part or part == b"--":
            continue
        head, _, content = part.partition(b"\r\n\r\n")
        headers = head.decode("utf-8", errors="replace")
        m = re.search(r'content-disposition:.*?name="([^"]+)"(?:;\s*filename="([^"]*)")?', headers, re.I)
        if not m:
            continue
        out[m.group(1)] = (m.group(2), content)
    return out


# ---------- HTTP ----------

class Handler(BaseHTTPRequestHandler):
    protocol_version = "HTTP/1.1"

    def log_message(self, format, *args):  # quiet-ish: one line per request to stderr
        sys.stderr.write("sidecar %s - %s\n" % (self.address_string(), format % args))

    def _send(self, code: int, payload: dict):
        body = json.dumps(payload).encode()
        self.send_response(code)
        self.send_header("content-type", "application/json")
        self.send_header("content-length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def _read_body(self) -> bytes:
        n = int(self.headers.get("content-length", 0))
        if n > MAX_UPLOAD:
            raise ValueError("upload too large")
        return self.rfile.read(n)

    def do_GET(self):
        try:
            if self.path == "/health":
                conn = db()
                s = conn.execute("SELECT count(*) c FROM sources").fetchone()["c"]
                ch = conn.execute("SELECT count(*) c FROM chunks").fetchone()["c"]
                return self._send(200, {"ok": True, "sources": s, "chunks": ch})
            if self.path == "/sources":
                conn = db()
                rows = conn.execute("SELECT * FROM sources ORDER BY created_at DESC").fetchall()
                return self._send(200, {"sources": [source_row(r) for r in rows]})
            if self.path.startswith("/search"):
                from urllib.parse import urlparse, parse_qs

                q = parse_qs(urlparse(self.path).query)
                query = (q.get("q", [""])[0] or "").strip()
                limit = min(int(q.get("limit", ["8"])[0]), 25)
                agent_id = (q.get("agent_id", [""])[0] or "").strip()
                if not query:
                    return self._send(400, {"error": "q is required"})
                conn = db()
                # FTS5 MATCH; quote each term to tolerate arbitrary input
                match = " OR ".join('"' + t.replace('"', "") + '"' for t in query.split() if t)
                if not match:
                    return self._send(400, {"error": "q has no searchable terms"})
                # Scope enforcement: an agent query never sees private sources,
                # and agent-scoped sources only surface for named agents.
                scope_sql = ""
                params: list = [match]
                if agent_id:
                    scope_sql = (
                        "AND (s.scope = 'workspace' OR (s.scope = 'agent' AND EXISTS "
                        "(SELECT 1 FROM json_each(s.allowed_agent_ids) je WHERE je.value = ?)))"
                    )
                    params.append(agent_id)
                params.append(limit)
                rows = conn.execute(
                    f"""SELECT c.chunk_id, c.source_id, c.chunk_index,
                              snippet(chunks, 3, '«', '»', '…', 32) AS snip,
                              bm25(chunks) AS score, s.name, s.citation_enabled
                       FROM chunks c JOIN sources s ON s.id = c.source_id
                       WHERE chunks MATCH ? AND s.indexing_status = 'ready' {scope_sql}
                       ORDER BY score LIMIT ?""",
                    params,
                ).fetchall()
                return self._send(200, {"results": [
                    {
                        "chunkId": r["chunk_id"],
                        "sourceId": r["source_id"],
                        "sourceName": r["name"],
                        "chunkIndex": r["chunk_index"],
                        "snippet": r["snip"],
                        "score": round(-r["score"], 4),
                        "citationEnabled": bool(r["citation_enabled"]),
                    }
                    for r in rows
                ]})
            m = re.fullmatch(r"/chunks/([\w:-]+)", self.path)
            if m:
                conn = db()
                r = conn.execute(
                    """SELECT c.chunk_id, c.source_id, c.chunk_index, c.text,
                              s.name, s.scope, s.citation_enabled, s.uri
                       FROM chunks c JOIN sources s ON s.id = c.source_id
                       WHERE c.chunk_id = ?""",
                    (m.group(1),),
                ).fetchone()
                if not r:
                    return self._send(404, {"error": "chunk not found"})
                return self._send(200, {
                    "chunkId": r["chunk_id"],
                    "sourceId": r["source_id"],
                    "sourceName": r["name"],
                    "sourceUri": r["uri"],
                    "scope": r["scope"],
                    "chunkIndex": r["chunk_index"],
                    "text": r["text"],
                    "citationEnabled": bool(r["citation_enabled"]),
                })
            return self._send(404, {"error": "not found"})
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": str(e)})

    def do_POST(self):
        try:
            if self.path == "/sources/upload":
                ctype = self.headers.get("content-type", "")
                m = re.search(r"boundary=([^;]+)", ctype)
                if "multipart/form-data" not in ctype or not m:
                    return self._send(400, {"error": "multipart/form-data required"})
                parts = parse_multipart(self._read_body(), m.group(1).strip('"'))
                if "file" not in parts or not parts["file"][0]:
                    return self._send(400, {"error": "file field required"})
                filename = str(parts["file"][0])
                content = parts["file"][1]
                name_field = parts.get("name")
                name = (name_field[1].decode(errors="replace").strip() if name_field else "") or os.path.basename(filename)
                sid = f"k-{uuid.uuid4().hex[:8]}"
                stored = FILES / f"{sid}{Path(filename).suffix.lower()[:12]}"
                stored.write_bytes(content)
                conn = db()
                conn.execute(
                    """INSERT INTO sources (id, type, name, uri, stored_path, scope, allowed_agent_ids,
                                            citation_enabled, indexing_status, created_at)
                       VALUES (?,?,?,?,?,?,?,?, 'pending', ?)""",
                    (
                        sid, "file", name, None, str(stored),
                        parts.get("scope", (None, b"private"))[1].decode() or "private",
                        parts.get("allowedAgentIds", (None, b""))[1].decode() or None,
                        1 if parts.get("citationEnabled", (None, b"true"))[1].decode() != "false" else 0,
                        now_iso(),
                    ),
                )
                conn.commit()
                count, err = index_source(conn, sid)
                row = conn.execute("SELECT * FROM sources WHERE id = ?", (sid,)).fetchone()
                out = source_row(row)
                out["chunkCount"] = count
                if err:
                    out["indexError"] = err
                return self._send(200, {"source": out})

            if self.path == "/sources/url":
                payload = json.loads(self._read_body() or b"{}")
                url = (payload.get("url") or "").strip()
                if not url.startswith(("http://", "https://")):
                    return self._send(400, {"error": "http(s) url required"})
                body, ctype = fetch_url(url)
                sid = f"k-{uuid.uuid4().hex[:8]}"
                ext = ".html" if "html" in ctype else ".pdf" if "pdf" in ctype else ".txt"
                stored = FILES / f"{sid}{ext}"
                stored.write_bytes(body)
                name = payload.get("name") or url
                conn = db()
                conn.execute(
                    """INSERT INTO sources (id, type, name, uri, stored_path, scope, allowed_agent_ids,
                                            citation_enabled, indexing_status, created_at)
                       VALUES (?,?,?,?,?,?,?,?, 'pending', ?)""",
                    (
                        sid, "url", name, url, str(stored),
                        payload.get("scope", "private"),
                        json.dumps(payload["allowedAgentIds"]) if payload.get("allowedAgentIds") else None,
                        0 if payload.get("citationEnabled") is False else 1,
                        now_iso(),
                    ),
                )
                conn.commit()
                count, err = index_source(conn, sid)
                row = conn.execute("SELECT * FROM sources WHERE id = ?", (sid,)).fetchone()
                out = source_row(row)
                out["chunkCount"] = count
                if err:
                    out["indexError"] = err
                return self._send(200, {"source": out})

            m = re.fullmatch(r"/sources/([\w-]+)/reindex", self.path)
            if m:
                conn = db()
                count, err = index_source(conn, m.group(1))
                if err == "source not found":
                    return self._send(404, {"error": err})
                row = conn.execute("SELECT * FROM sources WHERE id = ?", (m.group(1),)).fetchone()
                out = source_row(row)
                out["chunkCount"] = count
                return self._send(200, {"source": out})

            return self._send(404, {"error": "not found"})
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": str(e)})

    def do_DELETE(self):
        try:
            m = re.fullmatch(r"/sources/([\w-]+)", self.path)
            if not m:
                return self._send(404, {"error": "not found"})
            conn = db()
            row = conn.execute("SELECT stored_path FROM sources WHERE id = ?", (m.group(1),)).fetchone()
            if not row:
                return self._send(404, {"error": "source not found"})
            conn.execute("DELETE FROM chunks WHERE source_id = ?", (m.group(1),))
            conn.execute("DELETE FROM sources WHERE id = ?", (m.group(1),))
            conn.commit()
            try:
                Path(row["stored_path"]).unlink(missing_ok=True)
            except OSError:
                pass
            return self._send(200, {"ok": True})
        except Exception as e:
            traceback.print_exc()
            return self._send(500, {"error": str(e)})


def main():
    DATA.mkdir(exist_ok=True)
    FILES.mkdir(exist_ok=True)
    db().close()  # ensure schema
    server = ThreadingHTTPServer((HOST, PORT), Handler)
    print(f"eaios-knowledge sidecar listening on http://{HOST}:{PORT}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()
